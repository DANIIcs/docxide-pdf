"""Office Math: inline and display equations in Calibri 11 text — fractions,
scripts, radicals, sums and integrals, delimiters, a function, a matrix and a
nested fraction, with a long inline zone that breaks after an operator.

usage: python3 generate.py [output.docx]
"""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml

MNS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


def r(t, sty=None):
    s = f'<m:rPr><m:sty m:val="{sty}"/></m:rPr>' if sty else ''
    t = t.replace('&', '&amp;').replace('<', '&lt;')
    return f'<m:r>{s}<m:t xml:space="preserve">{t}</m:t></m:r>'


def f(n, d):
    return f'<m:f><m:num>{n}</m:num><m:den>{d}</m:den></m:f>'


def sup(b, s):
    return f'<m:sSup><m:e>{b}</m:e><m:sup>{s}</m:sup></m:sSup>'


def sub(b, s):
    return f'<m:sSub><m:e>{b}</m:e><m:sub>{s}</m:sub></m:sSub>'


def rad(e):
    return f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{e}</m:e></m:rad>'


def nary(ch, lo, hi, e):
    return (f'<m:nary><m:naryPr><m:chr m:val="{ch}"/></m:naryPr><m:sub>{lo}</m:sub>'
            f'<m:sup>{hi}</m:sup><m:e>{e}</m:e></m:nary>')


def d(e, beg=None, end=None):
    pr = (f'<m:begChr m:val="{beg}"/>' if beg else '') + (f'<m:endChr m:val="{end}"/>' if end else '')
    return f'<m:d>{"<m:dPr>" + pr + "</m:dPr>" if pr else ""}<m:e>{e}</m:e></m:d>'


def func(name, e):
    return f'<m:func><m:fName>{r(name, "p")}</m:fName><m:e>{e}</m:e></m:func>'


def matrix(rows):
    mr = ''.join('<m:mr>' + ''.join(f'<m:e>{c}</m:e>' for c in row) + '</m:mr>' for row in rows)
    return f'<m:m>{mr}</m:m>'


def inline(x):
    return f'<m:oMath {MNS}>{x}</m:oMath>'


def display(x):
    return f'<m:oMathPara {MNS}><m:oMath>{x}</m:oMath></m:oMathPara>'


def para(doc, parts):
    p = doc.add_paragraph()
    for part in parts:
        if part.startswith('<m:'):
            p._p.append(parse_xml(part))
        else:
            p.add_run(part)


def main(out):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    doc.add_heading('Office Math layout', level=1)
    para(doc, ['The period of a pendulum is ', inline(r('T=2π') + rad(f(r('l'), r('g')))),
               ' for small angles ', inline(r('θ≪') + f(r('π'), r('2'))),
               ', and the energy ', inline(r('E=m') + sup(r('c'), r('2'))), ' does not depend on ',
               inline(sub(r('x'), r('i'))), ' or ', inline(sup(r('y'), r('n+1'))), '.'])
    para(doc, ['The roots of a quadratic are'])
    para(doc, [display(r('x=') + f(r('-b±') + rad(sup(r('b'), r('2')) + r('-4ac')), r('2a')))])
    para(doc, ['and a sum and an integral read'])
    para(doc, [display(nary('∑', r('i=1'), r('n'), sup(r('i'), r('2'))) + r('=')
                       + f(r('n') + d(r('n+1')) + d(r('2n+1')), r('6')))])
    para(doc, [display(nary('∫', r('0'), r('1'), sup(r('x'), r('2')) + r('dx')) + r('=')
                       + f(r('1'), r('3')))])
    para(doc, ['Inline, the same sum is ', inline(nary('∑', r('i=1'), r('n'), r('i'))),
               ' and a function ', inline(func('sin', d(r('x+y')))), ' with ',
               inline(d(f(r('a'), r('b')))), ' in brackets. A long inline expression '
               'breaks after an operator: ',
               inline(r('f') + d(r('x')) + r('=') + sub(r('a'), r('0')) + r('+')
                      + sub(r('a'), r('1')) + r('x+') + sub(r('a'), r('2')) + sup(r('x'), r('2'))
                      + r('+⋯+') + sub(r('a'), r('n')) + sup(r('x'), r('n'))),
               ' and the text continues after it.'])
    para(doc, ['A rotation matrix and a nested fraction:'])
    para(doc, [display(r('R=') + d(matrix([[func('cos', r('θ')), r('-') + func('sin', r('θ'))],
                                         [func('sin', r('θ')), func('cos', r('θ'))]]), '[', ']')
                       + r(',  ') + f(f(r('a'), r('b')), r('c')))])
    doc.save(out)


if __name__ == '__main__':
    main(sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).with_name('input.docx')))
