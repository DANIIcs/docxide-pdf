"""Right-to-left paragraphs: w:bidi alignment and indents, w:rtl runs with
Latin, numbers and brackets, a bidi numbered list, Hebrew, and plain-run
Arabic with comma clusters. Arial 12 throughout.

usage: python3 generate.py [output.docx]
"""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AR = "يعتبر الأرز من أهم الحبوب التي يحتاجها جسم الإنسان"
AR_LONG = ("يعتبر الأرز من أهم الحبوب التي يحتاجها جسم الإنسان ومن مميزاته أنه يقدم في جميع "
           "الفصول وتكاد لا تخلو مائدة منه يتناوله الكبار والصغار والمرضى والأصحاء وهناك حوالي "
           "مئة وخمسون صنفا من الأرز في العالم كله")


def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def para(doc, runs, bidi=True, jc=None, ind=None, style=None):
    """runs: (text, rtl) pairs."""
    p = doc.add_paragraph(style=style)
    ppr = p._p.get_or_add_pPr()
    if bidi:
        ppr.append(el("w:bidi"))
    if ind:
        ppr.append(el("w:ind", **{f"w:{k}": v for k, v in ind.items()}))
    if jc:
        ppr.append(el("w:jc", **{"w:val": jc}))
    for text, rtl in runs:
        run = p.add_run(text)
        if rtl:
            run._r.get_or_add_rPr().append(el("w:rtl"))
    return p


def generate(path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(612), Pt(792)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Pt(72))
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = el("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Arial")
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    rpr.append(el("w:szCs", **{"w:val": 24}))
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.0

    # Word's own pattern: Arabic and its spaces in rtl runs, Latin plain.
    para(doc, [("النص ", True), ("Hello World", False), (" مع 12 نهاية.", True)])
    para(doc, [("النص (قوس) [مربع] و 4.5% نهاية", True)])
    para(doc, [(AR_LONG, True)], jc="both", ind={"left": 720, "right": 1440, "firstLine": 360})
    para(doc, [(AR, True)], jc="center")
    para(doc, [(AR, True)], jc="left")
    para(doc, [(AR, True)], jc="right")
    for i in range(2):
        para(doc, [(AR + " " + str(i + 1), True)], style="List Number")
    para(doc, [("שלום עולם ומלואו.", True)])
    # A plain LTR paragraph of Arabic, as documents without markup have it.
    para(doc, [("منه, يتناوله الكبار والصغار – مواد دهنية حوالي 150 صنف من الأرز .", False)],
         bidi=False)
    doc.save(path)


if __name__ == "__main__":
    generate(sys.argv[1] if len(sys.argv) > 1 else str(Path(__file__).with_name("input.docx")))
